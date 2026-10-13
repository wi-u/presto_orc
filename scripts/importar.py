#!/usr/bin/env python3
"""Converte a planilha de preços brutos da PRESTO em data/precos.json.

Uso:  python scripts/importar.py   (usa a única planilha .xlsx da pasta dados/)

- Lê as colunas pelo NOME do cabeçalho (não pela posição).
- Valida tudo antes de gravar. Se houver ERRO, nada é gravado.
- Mostra um relatório do que mudou em relação ao precos.json anterior.
"""
import csv, json, sys, datetime, re
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
def _achar_xlsx():
    if len(sys.argv) > 1:
        return Path(sys.argv[1])
    arqs = sorted((RAIZ / "dados").glob("*.xlsx"))
    if len(arqs) != 1:
        sys.exit(f"ERRO: a pasta dados/ deve ter exatamente 1 planilha .xlsx (achei {len(arqs)}). "
                 "Apague a antiga ao subir a nova.")
    return arqs[0]
XLSX = _achar_xlsx()
SAIDA = RAIZ / "data" / "precos.json"
AJUSTES = RAIZ / "data" / "ajustes_itens.csv"
POLITICA = RAIZ / "data" / "politica.json"
RELATORIO = RAIZ / "data" / "ultimo_relatorio.txt"

COLUNAS = {  # nome no JSON -> como o cabeçalho começa na planilha
    "cod": "COD", "descricao": "DESCRI", "mult_vda": "MULT", "preco_bruto": "PRE",
    "ipi": "IPI", "ncm": "NCM", "origem": "ORIGEM", "cst": "CST", "cest": "CEST", "st": "ST",
}

erros, avisos, infos = [], [], []   # infos = casos já conhecidos e aceitos

def norm(v):
    return "" if v is None else str(v).strip()

def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    t = norm(v).replace("R$", "").replace(" ", "")
    if not t:
        return None
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None

def achar_cabecalho(ws):
    for r in ws.iter_rows(values_only=True):
        if norm(r[0]).upper() == "COD":
            mapa = {}
            for i, c in enumerate(r):
                t = norm(c).upper()
                for chave, inicio in COLUNAS.items():
                    if chave not in mapa and t.startswith(inicio):
                        mapa[chave] = i
            faltando = [k for k in ("cod", "descricao", "mult_vda", "preco_bruto", "ipi", "ncm", "origem") if k not in mapa]
            if faltando:
                sys.exit(f"ERRO: colunas não encontradas no cabeçalho: {faltando}")
            return mapa
    sys.exit("ERRO: não achei a linha de cabeçalho (primeira coluna 'COD').")

def ler_ajustes():
    aj = {}
    if AJUSTES.exists():
        with open(AJUSTES, encoding="utf-8") as f:
            for lin in csv.DictReader(f):
                if norm(lin.get("cod")):
                    aj[norm(lin["cod"])] = norm(lin.get("linha"))
    return aj

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    mapa = achar_cabecalho(ws)
    ajustes = ler_ajustes()
    linhas_validas = set(json.load(open(POLITICA, encoding="utf-8"))["linhas"])
    for cod, linha in ajustes.items():
        if linha not in linhas_validas:
            erros.append(f"ajustes_itens.csv: linha '{linha}' (cod {cod}) não existe em politica.json")

    itens, vistos, categoria = [], {}, ""
    for n, r in enumerate(ws.iter_rows(values_only=True), start=1):
        g = lambda k: r[mapa[k]] if k in mapa and mapa[k] < len(r) else None
        cod = norm(g("cod"))
        if not cod or cod.upper() == "COD":
            continue
        desc = norm(g("descricao"))
        if not desc:                       # linha de categoria (ex.: "BANCADAS")
            categoria = cod.title()
            continue
        preco = num(g("preco_bruto"))
        mult = num(g("mult_vda"))
        ipi = num(g("ipi"))
        ncm = norm(g("ncm"))
        origem = norm(g("origem"))
        if origem.endswith(".0"):
            origem = origem[:-2]
        sem_preco = norm(g("preco_bruto")) == ""
        if sem_preco:                      # ex.: expositores "ver política"
            preco, ipi = 0.0, 0.0 if ipi is None else ipi
        if preco is None or preco < 0:
            erros.append(f"linha {n} ({cod}): preço inválido: {g('preco_bruto')!r}"); continue
        if mult is None or mult < 1 or mult != int(mult):
            txt = norm(g('mult_vda'))
            if re.fullmatch(r"\d+\s*(prat|col)\w*", txt, re.I):   # "2 prat", "4 col": o kit já traz isso
                infos.append(f"{cod}: múltiplo '{txt}' tratado como 1")
            else:
                avisos.append(f"{cod}: múltiplo '{txt}' não é um número — usei 1")
            mult = 1
        if ipi is None or not 0 <= ipi <= 0.5:
            erros.append(f"linha {n} ({cod}): IPI inválido: {g('ipi')!r} (esperado decimal, ex.: 0.0325)"); continue
        if preco == 0:
            if ncm == "-":   # expositor/bônus: fica sem preço de propósito
                infos.append(f"{cod}: expositor sem preço (bônus)")
            else:
                avisos.append(f"{cod} - {desc}: preço zero")
        elif origem not in ("0", "1"):
            erros.append(f"linha {n} ({cod}): origem inválida: {origem!r} (esperado 0 ou 1)")
            continue
        if preco > 0 and not ncm:
            erros.append(f"linha {n} ({cod}): NCM vazio")
        if cod in vistos:
            ant_i = vistos[cod]
            if ant_i["descricao"] == re.sub(r"\s+", " ", desc) and ant_i["preco_bruto"] == round(preco, 2):
                infos.append(f"linha {n}: {cod} repetido com os mesmos dados — contado uma vez")
            else:
                erros.append(f"linha {n}: código {cod} repetido com dados DIFERENTES (preço/descrição)")
            continue
        item = {
            "cod": cod, "descricao": re.sub(r"\s+", " ", desc), "und": "UN",
            "mult_vda": int(mult), "preco_bruto": round(preco, 2), "ipi": ipi,
            "ncm": ncm, "origem": origem, "importado": origem == "1",
            "cst": norm(g("cst")), "cest": norm(g("cest")), "st": norm(g("st")),
            "categoria": categoria,
        }
        if ajustes.get(cod):
            item["linha"] = ajustes[cod]
        vistos[cod] = item
        itens.append(item)

    for cod in ajustes:
        if cod not in vistos:
            avisos.append(f"ajustes_itens.csv: código {cod} não existe na planilha")
    if len(itens) < 100:
        erros.append(f"só {len(itens)} itens lidos — a planilha parece incompleta")

    # relatório de mudanças vs. versão anterior
    rel = []
    if SAIDA.exists():
        ant = {i["cod"]: i for i in json.load(open(SAIDA, encoding="utf-8"))["itens"]}
        novo = {i["cod"]: i for i in itens}
        add = sorted(set(novo) - set(ant)); rem = sorted(set(ant) - set(novo))
        mud = [(c, ant[c]["preco_bruto"], novo[c]["preco_bruto"]) for c in novo if c in ant and ant[c]["preco_bruto"] != novo[c]["preco_bruto"]]
        rel.append(f"Itens: {len(ant)} -> {len(novo)} | novos: {len(add)} | removidos: {len(rem)} | preço alterado: {len(mud)}")
        for c in add[:30]: rel.append(f"  + {c} {novo[c]['descricao']} (R$ {novo[c]['preco_bruto']:.2f})")
        for c in rem[:30]: rel.append(f"  - {c} {ant[c]['descricao']}")
        grandes = sorted(((abs(b / a - 1) if a else 1, c, a, b) for c, a, b in mud), reverse=True)
        for p, c, a, b in grandes[:15]:
            rel.append(f"  ~ {c}: R$ {a:.2f} -> R$ {b:.2f} ({(b / a - 1) * 100:+.1f}%)" if a else f"  ~ {c}: R$ {a:.2f} -> R$ {b:.2f}")
        if any(p > 0.5 for p, *_ in grandes):
            avisos.append("há itens com variação de preço acima de 50% — confira o relatório")
    else:
        rel.append(f"Primeira importação: {len(itens)} itens.")

    texto = ["RELATÓRIO DE IMPORTAÇÃO — " + datetime.datetime.now().strftime("%d/%m/%Y %H:%M"),
             f"Arquivo: {XLSX.name}", ""] + rel
    texto += ["", f"Já conhecidos e aceitos ({len(infos)}): " + "; ".join(infos) if infos else ""]
    texto += ["", f"AVISOS ({len(avisos)}):"] + [" - " + a for a in avisos]
    texto += ["", f"ERROS ({len(erros)}):"] + [" - " + e for e in erros]
    print("\n".join(texto))
    if erros:
        sys.exit("\nImportação CANCELADA: corrija os erros acima. Nada foi gravado.")

    nome = XLSX.stem
    m = re.search(r"([A-Z]{3}\d{2})", nome.upper())
    dados = {
        "tabela": m.group(1) if m else nome,
        "gerado_em": datetime.datetime.now().strftime("%Y-%m-%d"),
        "total_itens": len(itens),
        "itens": itens,
    }
    SAIDA.write_text(json.dumps(dados, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    RELATORIO.write_text("\n".join(texto) + "\n", encoding="utf-8")
    print(f"\nOK: {len(itens)} itens gravados em {SAIDA.relative_to(RAIZ)}")

main()
